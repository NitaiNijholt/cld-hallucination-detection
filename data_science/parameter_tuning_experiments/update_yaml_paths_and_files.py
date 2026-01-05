#!/usr/bin/env python3
"""
1. Copy prompts_citation_mechanistic_lit.yaml -> prompts_citation_mechanistic_original.yaml
2. Update yaml_path in mechanistic_original Excel files only
"""

import pandas as pd
from pathlib import Path
import openpyxl
import shutil

def copy_yaml_file():
    """Copy mechanistic_lit.yaml to mechanistic_original.yaml."""
    prompts_dir = Path("/home/nitai/code/causalix.ai/data_science/parameter_tuning_experiments/alternative_prompts")
    
    print("="*80)
    print("STEP 1: COPYING YAML FILE")
    print("="*80)
    
    # Copy mechanistic_lit.yaml -> mechanistic_original.yaml
    source = prompts_dir / "prompts_citation_mechanistic_lit.yaml"
    target = prompts_dir / "prompts_citation_mechanistic_original.yaml"
    
    if source.exists():
        if not target.exists():
            shutil.copy2(source, target)
            print(f"✓ Copied: {source.name} → {target.name}")
        else:
            print(f"⚠️  Already exists: {target.name}")
            print(f"   Checking if it matches source...")
            with open(source, 'r') as f1, open(target, 'r') as f2:
                if f1.read() == f2.read():
                    print(f"   ✓ Files are identical, no action needed")
                else:
                    print(f"   ⚠️  Files differ! Will overwrite...")
                    shutil.copy2(source, target)
                    print(f"   ✓ Overwrote with mechanistic_lit content")
    else:
        print(f"❌ Source not found: {source.name}")
    
    print()


def update_yaml_path_in_params(excel_path: Path, old_yaml_name: str, new_yaml_name: str):
    """Update the yaml_path in the Params sheet."""
    try:
        # Load workbook
        wb = openpyxl.load_workbook(excel_path)
        
        if 'Params' not in wb.sheetnames:
            print(f"  ⚠️  No Params sheet")
            return False
        
        ws = wb['Params']
        
        # Find and update the yaml_path row
        updated = False
        for row in ws.iter_rows(min_row=1):
            if row[0].value and 'yaml_path' in str(row[0].value).lower():
                old_path = str(row[1].value)
                if old_yaml_name in old_path:
                    new_path = old_path.replace(old_yaml_name, new_yaml_name)
                    row[1].value = new_path
                    updated = True
                    print(f"  ✓ Updated yaml_path")
                    break
        
        if updated:
            wb.save(excel_path)
            return True
        else:
            print(f"  ⚠️  yaml_path doesn't contain '{old_yaml_name}'")
            return False
            
    except Exception as e:
        print(f"  ❌ Error: {e}")
        return False


def update_excel_files():
    """Update yaml_path in mechanistic_original Excel files only."""
    base_dir = Path("/home/nitai/code/causalix.ai/final_runs/RQ1a_gt_lit_citation")
    
    print("="*80)
    print("STEP 2: UPDATING YAML PATHS IN MECHANISTIC_ORIGINAL EXCEL FILES")
    print("="*80)
    
    print("\n📄 Updating mechanistic_original files:")
    print("   prompts_citation_mechanistic.yaml → prompts_citation_mechanistic_original.yaml")
    print("-"*80)
    
    mechanistic_original_files = [f for f in base_dir.rglob("*_mechanistic_original_*.xlsx") 
                                   if 'deprecated' not in str(f) and '.backup' not in str(f)]
    
    count_updated = 0
    for excel_path in sorted(mechanistic_original_files):
        print(f"{excel_path.relative_to(base_dir)}")
        if update_yaml_path_in_params(excel_path, "prompts_citation_mechanistic.yaml", 
                                       "prompts_citation_mechanistic_original.yaml"):
            count_updated += 1
    
    print(f"\n✓ Updated {count_updated}/{len(mechanistic_original_files)} files\n")


def main():
    print("="*80)
    print("UPDATING YAML FILE AND MECHANISTIC_ORIGINAL EXCEL PARAMS")
    print("="*80)
    print()
    
    # Step 1: Copy yaml file
    copy_yaml_file()
    
    # Step 2: Update mechanistic_original Excel files only
    update_excel_files()
    
    print("="*80)
    print("✅ ALL UPDATES COMPLETE")
    print("="*80)
    print("\nNote: mechanistic files (former mechanistic_lit) were not modified.")


if __name__ == "__main__":
    main()
