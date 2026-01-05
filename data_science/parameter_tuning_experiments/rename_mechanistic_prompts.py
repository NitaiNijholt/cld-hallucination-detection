#!/usr/bin/env python3
"""
Rename mechanistic prompts in RQ1a_gt_lit_citation:
1. mechanistic -> mechanistic_original (files and Params sheet)
2. mechanistic_lit -> mechanistic (files and Params sheet)
"""

import pandas as pd
from pathlib import Path
import shutil
import openpyxl
import re

def update_excel_params_sheet(excel_path: Path, old_prompt: str, new_prompt: str):
    """Update the prompt name in the Params sheet of an Excel file."""
    try:
        # Load workbook
        wb = openpyxl.load_workbook(excel_path)
        
        if 'Params' not in wb.sheetnames:
            print(f"  ⚠️  No Params sheet in {excel_path.name}")
            return False
        
        ws = wb['Params']
        
        # Find and update the prompt row
        updated = False
        for row in ws.iter_rows(min_row=1):
            if row[0].value and 'citation_strategy' in str(row[0].value):
                if row[1].value == old_prompt:
                    row[1].value = new_prompt
                    updated = True
                    break
        
        if updated:
            wb.save(excel_path)
            print(f"  ✓ Updated Params sheet: {old_prompt} → {new_prompt}")
            return True
        else:
            print(f"  ⚠️  Could not find {old_prompt} in Params sheet")
            return False
            
    except Exception as e:
        print(f"  ❌ Error updating {excel_path.name}: {e}")
        return False


def rename_files_and_update_params(base_dir: Path, old_pattern: str, new_pattern: str, old_prompt: str, new_prompt: str):
    """Rename files matching pattern and update their Params sheets."""
    
    files_to_rename = list(base_dir.rglob(old_pattern))
    
    if not files_to_rename:
        print(f"No files found matching: {old_pattern}")
        return
    
    print(f"\n{'='*80}")
    print(f"Processing: {old_pattern} → {new_pattern}")
    print(f"Found {len(files_to_rename)} files")
    print(f"{'='*80}\n")
    
    for old_path in sorted(files_to_rename):
        print(f"📄 {old_path.relative_to(base_dir)}")
        
        # Step 1: Update Params sheet
        update_excel_params_sheet(old_path, old_prompt, new_prompt)
        
        # Step 2: Rename file
        new_filename = old_path.name.replace(old_pattern.replace('*', ''), new_pattern.replace('*', ''))
        new_path = old_path.parent / new_filename
        
        if new_path.exists():
            print(f"  ⚠️  Target already exists: {new_filename}")
            continue
        
        try:
            shutil.move(str(old_path), str(new_path))
            print(f"  ✓ Renamed file: {new_filename}\n")
        except Exception as e:
            print(f"  ❌ Error renaming: {e}\n")


def main():
    base_dir = Path("/home/nitai/code/causalix.ai/final_runs/RQ1a_gt_lit_citation")
    
    if not base_dir.exists():
        print(f"Error: Base directory not found: {base_dir}")
        return
    
    print("="*80)
    print("RENAMING MECHANISTIC PROMPTS IN RQ1a_gt_lit_citation")
    print("="*80)
    print(f"Base directory: {base_dir}\n")
    
    # Step 1: Rename mechanistic → mechanistic_original FIRST (to avoid naming collision)
    print("\n🔄 STEP 1: mechanistic → mechanistic_original")
    # Find files that are mechanistic but NOT mechanistic_lit
    all_mechanistic = set(base_dir.rglob("*_mechanistic_*.xlsx"))
    mechanistic_lit = set(base_dir.rglob("*_mechanistic_lit_*.xlsx"))
    mechanistic_only = all_mechanistic - mechanistic_lit
    
    if mechanistic_only:
        print(f"Found {len(mechanistic_only)} mechanistic (non-lit) files")
        for old_path in sorted(mechanistic_only):
            print(f"📄 {old_path.relative_to(base_dir)}")
            
            # Update Params sheet
            update_excel_params_sheet(old_path, "mechanistic", "mechanistic_original")
            
            # Rename file by inserting _original between mechanistic and the timestamp
            # Pattern: judged_*_mechanistic_TIMESTAMP.xlsx -> judged_*_mechanistic_original_TIMESTAMP.xlsx
            new_filename = re.sub(r'_mechanistic_(\d{8}_\d{6})', r'_mechanistic_original_\1', old_path.name)
            new_path = old_path.parent / new_filename
            
            if new_path.exists():
                print(f"  ⚠️  Target already exists: {new_filename}")
                continue
            
            try:
                shutil.move(str(old_path), str(new_path))
                print(f"  ✓ Renamed file: {new_filename}\n")
            except Exception as e:
                print(f"  ❌ Error renaming: {e}\n")
    
    # Step 2: Rename mechanistic_lit → mechanistic
    print("\n🔄 STEP 2: mechanistic_lit → mechanistic")
    mechanistic_lit_files = list(base_dir.rglob("*_mechanistic_lit_*.xlsx"))
    
    if mechanistic_lit_files:
        print(f"Found {len(mechanistic_lit_files)} mechanistic_lit files")
        for old_path in sorted(mechanistic_lit_files):
            print(f"📄 {old_path.relative_to(base_dir)}")
            
            # Update Params sheet
            update_excel_params_sheet(old_path, "mechanistic_lit", "mechanistic")
            
            # Rename file by removing _lit from mechanistic_lit
            # Pattern: judged_*_mechanistic_lit_TIMESTAMP.xlsx -> judged_*_mechanistic_TIMESTAMP.xlsx
            new_filename = re.sub(r'_mechanistic_lit_(\d{8}_\d{6})', r'_mechanistic_\1', old_path.name)
            new_path = old_path.parent / new_filename
            
            if new_path.exists():
                print(f"  ⚠️  Target already exists: {new_filename}")
                continue
            
            try:
                shutil.move(str(old_path), str(new_path))
                print(f"  ✓ Renamed file: {new_filename}\n")
            except Exception as e:
                print(f"  ❌ Error renaming: {e}\n")
    
    print("\n" + "="*80)
    print("✅ RENAMING COMPLETE")
    print("="*80)


if __name__ == "__main__":
    main()

