#!/usr/bin/env python3
"""
Update yaml_path in Params sheet of renamed Excel files.
"""

import pandas as pd
from pathlib import Path
import openpyxl

def update_yaml_path_in_params(excel_path: Path, old_yaml_substr: str, new_yaml_substr: str):
    """Update the yaml_path in the Params sheet."""
    try:
        # Load workbook
        wb = openpyxl.load_workbook(excel_path)
        
        if 'Params' not in wb.sheetnames:
            print(f"  ⚠️  No Params sheet in {excel_path.name}")
            return False
        
        ws = wb['Params']
        
        # Find and update the yaml_path row
        updated = False
        for row in ws.iter_rows(min_row=1):
            if row[0].value and 'yaml_path' in str(row[0].value).lower():
                old_path = str(row[1].value)
                if old_yaml_substr in old_path:
                    new_path = old_path.replace(old_yaml_substr, new_yaml_substr)
                    row[1].value = new_path
                    updated = True
                    print(f"  ✓ Updated yaml_path:")
                    print(f"    Old: {old_path}")
                    print(f"    New: {new_path}")
                    break
        
        if updated:
            wb.save(excel_path)
            return True
        else:
            print(f"  ⚠️  yaml_path not found or doesn't contain '{old_yaml_substr}'")
            return False
            
    except Exception as e:
        print(f"  ❌ Error updating {excel_path.name}: {e}")
        return False


def main():
    base_dir = Path("/home/nitai/code/causalix.ai/final_runs/RQ1a_gt_lit_citation")
    
    if not base_dir.exists():
        print(f"Error: Base directory not found: {base_dir}")
        return
    
    print("="*80)
    print("UPDATING YAML PATHS IN PARAMS SHEETS")
    print("="*80)
    print(f"Base directory: {base_dir}\n")
    
    # Step 1: Update mechanistic_original files (old mechanistic)
    print("\n🔄 STEP 1: Updating mechanistic_original files")
    print("yaml_path: prompts_citation_mechanistic.yaml → prompts_citation_mechanistic_original.yaml")
    print("="*80)
    
    mechanistic_original_files = list(base_dir.rglob("*_mechanistic_original_*.xlsx"))
    print(f"Found {len(mechanistic_original_files)} files\n")
    
    for excel_path in sorted(mechanistic_original_files):
        if 'deprecated' in str(excel_path) or '.backup' in str(excel_path):
            print(f"📄 {excel_path.relative_to(base_dir)} [SKIPPING - deprecated/backup]")
            continue
            
        print(f"📄 {excel_path.relative_to(base_dir)}")
        update_yaml_path_in_params(
            excel_path,
            old_yaml_substr="prompts_citation_mechanistic.yaml",
            new_yaml_substr="prompts_citation_mechanistic_original.yaml"
        )
        print()
    
    # Step 2: Update mechanistic files (old mechanistic_lit)
    print("\n🔄 STEP 2: Updating mechanistic files (from mechanistic_lit)")
    print("yaml_path: prompts_citation_mechanistic_lit.yaml → prompts_citation_mechanistic.yaml")
    print("="*80)
    
    # Find mechanistic files that are NOT mechanistic_original
    all_mechanistic = set(base_dir.rglob("*_mechanistic_*.xlsx"))
    mechanistic_original = set(base_dir.rglob("*_mechanistic_original_*.xlsx"))
    mechanistic_only = all_mechanistic - mechanistic_original
    
    print(f"Found {len(mechanistic_only)} files\n")
    
    for excel_path in sorted(mechanistic_only):
        if 'deprecated' in str(excel_path) or '.backup' in str(excel_path):
            print(f"📄 {excel_path.relative_to(base_dir)} [SKIPPING - deprecated/backup]")
            continue
            
        print(f"📄 {excel_path.relative_to(base_dir)}")
        update_yaml_path_in_params(
            excel_path,
            old_yaml_substr="prompts_citation_mechanistic_lit.yaml",
            new_yaml_substr="prompts_citation_mechanistic.yaml"
        )
        print()
    
    print("\n" + "="*80)
    print("✅ YAML PATH UPDATES COMPLETE")
    print("="*80)


if __name__ == "__main__":
    main()


